"""文档解析补强的回归测试：Word 表格 + PDF 页码标记。

- Word 之前只读 doc.paragraphs，表格内容完全丢失；现在按文中原有顺序把段落和
  表格都转出来，用真实 python-docx 生成的 .docx 文件跑（不是 mock，因为
  python-docx 本来就是项目依赖，用它自己造测试夹具最可靠）。
- PDF 页码用 mock 的 PdfReader（生成真实可提取文字的 PDF 需要 reportlab 这类
  没有的依赖），只验证每页文字前插入了 [第N页] 标记这件事本身。
"""
import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch

from service.rag import rag_service


class ParsePdfPageMarkersTest(unittest.TestCase):
    def test_inserts_page_markers_in_order(self):
        page1 = MagicMock()
        page1.extract_text.return_value = "第一页内容"
        page2 = MagicMock()
        page2.extract_text.return_value = "第二页内容"
        with patch("pypdf.PdfReader") as mock_reader_cls:
            mock_reader_cls.return_value.pages = [page1, page2]
            result = rag_service._parse_pdf("fake.pdf")

        self.assertIn("[第1页]\n第一页内容", result)
        self.assertIn("[第2页]\n第二页内容", result)
        # 页1 在页2 前面
        self.assertLess(result.index("[第1页]"), result.index("[第2页]"))

    def test_blank_pages_are_skipped_not_marked(self):
        blank_page = MagicMock()
        blank_page.extract_text.return_value = ""  # 扫描件/图片页常见：提取不出文字
        text_page = MagicMock()
        text_page.extract_text.return_value = "有内容的页"
        with patch("pypdf.PdfReader") as mock_reader_cls:
            mock_reader_cls.return_value.pages = [blank_page, text_page]
            result = rag_service._parse_pdf("fake.pdf")

        self.assertNotIn("[第1页]", result)  # 空白页不留占位标记
        self.assertIn("[第2页]\n有内容的页", result)


class ParseDocxTableTest(unittest.TestCase):
    def _build_docx(self, path: str):
        import docx
        doc = docx.Document()
        doc.add_paragraph("引言段落")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "套餐"
        table.cell(0, 1).text = "价格"
        table.cell(1, 0).text = "专业版"
        table.cell(1, 1).text = "39元/月"
        doc.add_paragraph("结尾段落")
        doc.save(path)

    def test_table_content_is_preserved_in_order(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = str(Path(tmp) / "sample.docx")
            self._build_docx(path)
            result = rag_service._parse_docx(path)

        self.assertIn("引言段落", result)
        self.assertIn("结尾段落", result)
        self.assertIn("【表格】", result)
        self.assertIn("套餐 | 价格", result)
        self.assertIn("专业版 | 39元/月", result)
        # 顺序：引言 → 表格 → 结尾，和文档里写入的顺序一致
        self.assertLess(result.index("引言段落"), result.index("【表格】"))
        self.assertLess(result.index("【表格】"), result.index("结尾段落"))

    def test_empty_table_rows_are_skipped(self):
        with tempfile.TemporaryDirectory() as tmp:
            import docx
            path = str(Path(tmp) / "empty_table.docx")
            doc = docx.Document()
            doc.add_paragraph("正文")
            doc.add_table(rows=1, cols=2)  # 全空表格，不该产生【表格】标记
            doc.save(path)

            result = rag_service._parse_docx(path)

        self.assertIn("正文", result)
        self.assertNotIn("【表格】", result)


if __name__ == "__main__":
    unittest.main()
